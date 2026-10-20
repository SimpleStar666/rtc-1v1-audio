#!/usr/bin/env python3
# 验证改后的简历：页数、残留 libwebrtc、内容完整性
import sys, io, subprocess, os
from docx import Document

path = r"C:\Users\小陶\Desktop\陶云飞-WebRTC音视频开发工程师.docx"
d = Document(path)

lines = []
lines.append('=== 全文（改后）===')
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t:
        lines.append(f'[{i}] {t}')

full = '\n'.join(lines)

# ---- 检查 1：是否还有 libwebrtc ----
bad = []
for i, p in enumerate(d.paragraphs):
    if 'libwebrtc' in p.text or 'libWebRTC' in p.text:
        bad.append((i, p.text[:60]))

lines.append('')
lines.append('=== 检查 1：残留 libwebrtc ===')
if bad:
    lines.append(f'  [FAIL] 还有 {len(bad)} 处残留：')
    for i, t in bad:
        lines.append(f'    [{i}] {t}')
else:
    lines.append('  [PASS] 已无 libwebrtc 残留')

# ---- 检查 2：视频关键词是否都进去了 ----
keywords = ['VP8', 'PSNR', '分片重组', '关键帧', 'PLI', 'YUV420']
lines.append('')
lines.append('=== 检查 2：视频关键词 ===')
for kw in keywords:
    n = full.count(kw)
    lines.append(f'  {kw:12s}: {n} 次 {"[OK]" if n > 0 else "[MISSING]"}')

# ---- 检查 3：自建组件是否写了 ----
selfbuild = ['libwebsockets', 'libnice', 'libsrtp2', 'OpenSSL', 'libvpx']
lines.append('')
lines.append('=== 检查 3：自建组件 ===')
for kw in selfbuild:
    n = full.count(kw)
    lines.append(f'  {kw:14s}: {n} 次 {"[OK]" if n > 0 else "[MISSING]"}')

# ---- 检查 4：实测数据是否写了 ----
metrics = ['29.8', '52', 'srflx', '100%/89%', '1100B', '300ms']
lines.append('')
lines.append('=== 检查 4：实测数据 ===')
for kw in metrics:
    n = full.count(kw)
    lines.append(f'  {kw:14s}: {n} 次 {"[OK]" if n > 0 else "[MISSING]"}')

# ---- 检查 5：页数（通过 PDF 转换后的页数判断）----
lines.append('')
lines.append('=== 检查 5：页数 ===')
lines.append('  （需要转 PDF 后统计，见下一步）')

# ---- 检查 6：段落总数（防止意外删改）----
lines.append('')
lines.append(f'=== 段落统计 ===')
lines.append(f'  总段落数: {len(d.paragraphs)}')
lines.append(f'  非空段落: {len([p for p in d.paragraphs if p.text.strip()])}')
lines.append(f'  表格数: {len(d.tables)}')

out = r"C:\Users\小陶\Documents\deepseek-harness\default-workspace\resume_verify.txt"
with io.open(out, 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines))
print(f'OK -> {out}')
