#!/usr/bin/env python3
# 导出简历全部内容（段落 + 表格），供核对
import sys
from docx import Document

path = sys.argv[1] if len(sys.argv) > 1 else '/mnt/c/Users/小陶/Documents/deepseek-harness/default-workspace/resume/陶云飞-WebRTC音视频开发工程师.docx'
d = Document(path)

print('=== 段落 ===')
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t:
        print(f'[{i}] {t}')

print('\n=== 表格 ===')
for ti, tbl in enumerate(d.tables):
    print(f'--- 表格 {ti} ({len(tbl.rows)} 行 x {len(tbl.columns)} 列) ---')
    for ri, row in enumerate(tbl.rows):
        cells = [c.text.strip().replace('\n', ' / ') for c in row.cells]
        print(f'  [{ri}] ' + ' | '.join(cells))

print('\n=== 节（页面设置）===')
for si, sec in enumerate(d.sections):
    print(f'  节{si}: {sec.page_width.cm:.2f} x {sec.page_height.cm:.2f} cm, '
          f'边距 上{sec.top_margin.cm:.2f} 下{sec.bottom_margin.cm:.2f} '
          f'左{sec.left_margin.cm:.2f} 右{sec.right_margin.cm:.2f}')
