#!/usr/bin/env python3
# 彻底去掉项目名称里的 libwebrtc
import io
from docx import Document

path = r"C:\Users\小陶\Desktop\陶云飞-WebRTC音视频开发工程师.docx"
d = Document(path)
p = d.paragraphs[38]
old = p.text.strip()

new = '项目名称：WebRTC 1v1 实时音视频通话系统（基于开源组件从零自建）'
if not p.runs:
    p.text = new
else:
    p.runs[0].text = new
    for r in p.runs[1:]:
        r.text = ''

d.save(path)
print(f'旧: {old}')
print(f'新: {new}')

# 复检全文
d2 = Document(path)
found = False
for i, pp in enumerate(d2.paragraphs):
    low = pp.text.lower()
    if 'libwebrtc' in low:
        print(f'  [FAIL] [{i}] {pp.text[:70]}')
        found = True
if not found:
    print('\n[PASS] 全文已无 libwebrtc 字样')

# 顺便检查 "语音通话" 是否还有残留（应该是音视频了）
print('\n=== 检查 "语音通话" 残留 ===')
found2 = False
for i, pp in enumerate(d2.paragraphs):
    if '语音通话' in pp.text:
        print(f'  [{i}] {pp.text[:70]}')
        found2 = True
if not found2:
    print('  [PASS] 已无"语音通话"（都是音视频）')
