#!/usr/bin/env python3
# 精简个人项目段落，压回 2 页（保留所有关键技术点和实测数据）
import io
from docx import Document

path = r"C:\Users\小陶\Desktop\陶云飞-WebRTC音视频开发工程师.docx"
d = Document(path)
paras = d.paragraphs

def set_text(idx, new_text, label):
    p = paras[idx]
    old_len = len(p.text.strip())
    if not p.runs:
        p.text = new_text
    else:
        p.runs[0].text = new_text
        for r in p.runs[1:]:
            r.text = ''
    print(f'  [{idx}] {label}: {old_len} -> {len(new_text)} 字')

# ============================================================
# 精简策略：
#   [40] 318 -> 约 150（项目描述只讲"做了什么"，细节留给职责）
#   [42] 153 -> 约 110
#   [43] 150 -> 约 105
#   [44] 126 -> 约 100
#   [45] 127 -> 约 100
#   [46] 186 -> 约 130（视频专项是亮点，保留最多）
#   [48] 163 -> 约 110
# ============================================================

# ---- [40] 项目描述 ----
set_text(40,
    '项目描述：从零自建 1v1 实时音视频通话系统（约 1.2 万行 C/C++，'
    '基于 libwebsockets / libnice / libsrtp2 / OpenSSL / libopus / libvpx 六个开源组件）：'
    '音视频编解码、自适应抖动缓冲、RTP 分片重组与关键帧请求、WebSocket 信令与 Trickle ICE、'
    'STUN 打洞、DTLS-SRTP 加密、NACK/FEC/GCC 弱网对抗。全链路有可复现实测数据与 44 项踩坑记录。',
    '项目描述')

# ---- [42] 编解码 ----
set_text(42,
    '1. 音视频编解码：Opus 实现 16kHz 单声道语音编解码（16kbps / SNR 29.8dB）；'
    'VP8(libvpx) 实现视频编解码（320×240@15fps / PSNR 52dB），'
    '处理 YUV420 stride 拷贝、CBR 码率控制、周期性关键帧与运行时码率调整。',
    '职责1-编解码')

# ---- [43] 信令 ----
set_text(43,
    '2. 信令服务：自建 WebSocket 信令服务器（libwebsockets），转发 offer / answer / ice-candidate；'
    '实现 Trickle ICE 边收集边检查以压缩首帧时间，协商时交换 ICE 凭证与 DTLS 证书指纹。',
    '职责2-信令')

# ---- [44] NAT/加密 ----
set_text(44,
    '3. NAT 穿透与加密：基于 libnice 做 ICE 候选收集与连通性检查，STUN 打洞获取公网映射（srflx）候选；'
    '媒体经 DTLS 握手派生密钥（RFC 5764）后走 SRTP 加密，并做双向证书指纹校验防中间人。',
    '职责3-NAT/加密')

# ---- [45] 弱网 ----
set_text(45,
    '4. 弱网对抗：NACK 重传 + XOR FEC 冗余（按丢包率动态调冗余率）+ GCC 延迟梯度与 AIMD 调码率；'
    '抖动缓冲目标延迟自适应，丢包 2%/5%/10%/20% 下恢复率实测 100%/89%/78%/52%。',
    '职责4-弱网')

# ---- [46] 视频专项（亮点，保留最多信息）----
set_text(46,
    '5. 视频链路专项：解决一帧超 MTU 必须分片（1100B/片 + 8 字节分片头），'
    '接收端按 frame_id 重组、缺片整帧丢弃避免花屏；针对"丢关键帧导致后续 P 帧全不可解码"，'
    '实现关键帧请求（PLI，限流 300ms）与周期性关键帧，丢包后 1 个 RTT 内自愈；'
    '并用统一指标做跨层自洽校验。',
    '职责5-视频专项')

# ---- [48] 自我评价 ----
set_text(48,
    '1. 具备从协议设计、缓冲调度、编解码到设备层落地的完整链路经验，能独立定位跨模块问题；'
    '理解延迟、抖动、卡顿之间的取舍关系，并用自建音视频 RTC 项目（含 44 项踩坑记录）验证过全链路；'
    '对 RTC 技术有持续投入，性格踏实，能与硬件、云端等多团队协作推进问题闭环。',
    '自我评价')

# ---- [9] 技能第一条也略缩（视频能力保留）----
set_text(9,
    '•  熟悉 WebRTC 音视频管线：APM 音频处理（AEC / ANS / AGC）、NetEq 抖动缓冲、ADM 音频设备层；'
    '视频侧熟悉 VP8 编解码、YUV420、RTP 分片重组与关键帧请求（PLI/FIR）；',
    '技能-音视频管线')

d.save(path)

# 统计
d2 = Document(path)
total = sum(len(p.text.strip()) for p in d2.paragraphs if p.text.strip())
print(f'\n总字符数: {total}（改前 3227）')
print('已保存')
