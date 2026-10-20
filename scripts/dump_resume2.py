#!/usr/bin/env python3
# 导出简历全部内容到 UTF-8 文件（避免 Windows 控制台编码问题）
import sys, io
from docx import Document

path = sys.argv[1]
out_path = sys.argv[2] if len(sys.argv) > 2 else 'resume_dump.txt'
d = Document(path)

lines = []
lines.append('=== 段落 ===')
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t:
        # 标记样式，便于判断标题层级
        st = p.style.name if p.style else ''
        lines.append(f'[{i}] ({st}) {t}')

lines.append('')
lines.append('=== 表格 ===')
for ti, tbl in enumerate(d.tables):
    lines.append(f'--- 表格 {ti} ({len(tbl.rows)} 行 x {len(tbl.columns)} 列) ---')
    for ri, row in enumerate(tbl.rows):
        cells = [c.text.strip().replace('\n', ' ⏎ ') for c in row.cells]
        lines.append(f'  [{ri}] ' + ' | '.join(cells))

lines.append('')
lines.append('=== 节（页面设置）===')
for si, sec in enumerate(d.sections):
    lines.append(f'  节{si}: {sec.page_width.cm:.2f} x {sec.page_height.cm:.2f} cm, '
                 f'边距 上{sec.top_margin.cm:.2f} 下{sec.bottom_margin.cm:.2f} '
                 f'左{sec.left_margin.cm:.2f} 右{sec.right_margin.cm:.2f}')

with io.open(out_path, 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines))
print(f'OK: {len(lines)} lines -> {out_path}')
