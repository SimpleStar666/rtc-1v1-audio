#!/usr/bin/env python3
# 分析溢出：找出哪些段落变长了，需要精简
import io
from docx import Document

path = r"C:\Users\小陶\Desktop\陶云飞-WebRTC音视频开发工程师.docx"
d = Document(path)

lines = []
lines.append('=== 段落长度（字符数），按长度降序 ===')
items = []
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t:
        items.append((len(t), i, t))

items.sort(reverse=True)
for n, i, t in items[:20]:
    lines.append(f'  [{i}] {n:4d} 字: {t[:75]}...')

lines.append('')
lines.append('=== 总字符数 ===')
lines.append(f'  {sum(n for n, _, _ in items)} 字符（非空段落）')

lines.append('')
lines.append('=== 需要精简的候选（个人项目部分，[38]-[46]）===')
for i in range(38, 47):
    t = d.paragraphs[i].text.strip()
    if t:
        lines.append(f'  [{i}] {len(t):4d} 字')

out = r"C:\Users\小陶\Documents\deepseek-harness\default-workspace\overflow_analysis.txt"
with io.open(out, 'w', encoding='utf-8') as f:
    f.write('\n'.join(lines))
print(f'OK -> {out}')
