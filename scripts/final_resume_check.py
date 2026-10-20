#!/usr/bin/env python3
# 最终核验：内容完整性 + 关键数据 + 页数 + 格式
import io, sys
from docx import Document

path = r"C:\Users\小陶\Desktop\陶云飞-WebRTC音视频开发工程师.docx"
d = Document(path)
paras = [p.text.strip() for p in d.paragraphs]
full = '\n'.join(paras)

L = []
def w(s=''): L.append(s)

w('=' * 60)
w('简历最终核验')
w('=' * 60)
w()

# ---- 1. 必须无 libwebrtc ----
w('【1】不应出现的关键词（会暴露与事实不符）')
for kw in ['libwebrtc', 'PeerConnection', 'chrome://webrtc-internals']:
    n = full.lower().count(kw.lower())
    w(f'  {kw:24s}: {n} 次  {"[FAIL] 需删除" if n else "[PASS]"}')
w()

# ---- 2. 必须有的视频关键词 ----
w('【2】视频能力关键词（音视频岗硬通货）')
video_kw = {
    'VP8': '视频编解码器',
    'libvpx': 'VP8 实现库',
    'PSNR': '视频质量指标',
    'YUV420': '视频像素格式',
    '分片': 'MTU 分片',
    '重组': '分片重组',
    '关键帧': '关键帧机制',
    'PLI': '关键帧请求',
}
for kw, desc in video_kw.items():
    n = full.count(kw)
    w(f'  {kw:10s} ({desc:12s}): {n:2d} 次  {"[OK]" if n else "[MISSING]"}')
w()

# ---- 3. 自建组件 ----
w('【3】自建组件（证明不是调 API）')
for kw in ['libwebsockets', 'libnice', 'libsrtp2', 'OpenSSL', 'libopus', 'libvpx']:
    n = full.count(kw)
    w(f'  {kw:14s}: {n} 次  {"[OK]" if n else "[MISSING]"}')
w()

# ---- 4. 实测数据 ----
w('【4】实测数据（面试可验证）')
for kw, desc in [('29.8', 'Opus SNR'), ('52', '视频 PSNR'), ('srflx', 'STUN 公网映射'),
                 ('100%/89%', '弱网恢复率'), ('1100B', '分片大小'),
                 ('300ms', 'PLI 限流'), ('44 项', '踩坑记录')]:
    n = full.count(kw)
    w(f'  {kw:12s} ({desc:12s}): {n} 次  {"[OK]" if n else "[MISSING]"}')
w()

# ---- 5. 工作经历完整性（不能改漏）----
w('【5】工作经历完整性')
checks = [
    ('中科创达', '第一家公司'),
    ('2025.4 ~ 2026.7', '中科创达时间'),
    ('武汉起点', '第二家公司'),
    ('2023.9 ~ 2025.3', '起点时间'),
    ('智能耳机实时音频通信链路', '中科创达项目名'),
    ('通信设备（军工）', '起点项目名'),
]
for kw, desc in checks:
    n = full.count(kw)
    w(f'  {desc:18s}: {"[OK]" if n else "[FAIL] 丢失!"}')
w()

# ---- 6. 中科创达 7 条职责都在 ----
w('【6】中科创达 7 条职责（一条都不能丢）')
duties = ['实时语音编解码', '应用层封包协议', '下行播放管线', '多角色音频路由',
          '音频设备层', '扩展应用', '提示音与音量管理']
for i, kw in enumerate(duties, 1):
    n = full.count(kw)
    w(f'  职责{i} {kw:16s}: {"[OK]" if n else "[FAIL] 丢失!"}')
w()

# ---- 7. 个人项目 5 条职责 ----
w('【7】个人项目 5 条职责')
proj_duties = ['音视频编解码', '信令服务', 'NAT 穿透与加密', '弱网对抗', '视频链路专项']
for i, kw in enumerate(proj_duties, 1):
    n = full.count(kw)
    w(f'  职责{i} {kw:16s}: {"[OK]" if n else "[MISSING]"}')
w()

# ---- 8. 段落统计 ----
w('【8】结构统计')
w(f'  总段落: {len(d.paragraphs)}')
w(f'  非空段落: {len([p for p in paras if p])}')
w(f'  表格: {len(d.tables)}')
w(f'  总字符: {sum(len(p) for p in paras if p)}')
w()

# ---- 9. 排版参数 ----
w('【9】排版（决定页数）')
s = d.sections[0]
w(f'  页面: {s.page_width.cm:.1f} x {s.page_height.cm:.1f} cm')
w(f'  边距: 上下 {s.top_margin.cm:.2f} / {s.bottom_margin.cm:.2f} cm, '
  f'左右 {s.left_margin.cm:.2f} / {s.right_margin.cm:.2f} cm')
ls = [p.paragraph_format.line_spacing for p in d.paragraphs
      if p.text.strip() and p.paragraph_format.line_spacing]
if ls:
    w(f'  行距: {min(ls):.2f} ~ {max(ls):.2f}')
w()

# ---- 10. 全文输出 ----
w('=' * 60)
w('全文（最终版）')
w('=' * 60)
for i, t in enumerate(paras):
    if t:
        w(f'[{i}] {t}')

out = r"C:\Users\小陶\Documents\deepseek-harness\default-workspace\resume_final_check.txt"
with io.open(out, 'w', encoding='utf-8') as f:
    f.write('\n'.join(L))
print(f'OK -> {out}')
